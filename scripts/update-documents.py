"""Replace unrelated template figures; align two supplied DOCX documents with the code.
Usage: python scripts/update-documents.py --reference <directory containing both originals>
"""
import argparse, re, json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parents[1]
args=argparse.ArgumentParser(); args.add_argument('--reference',type=Path,default=ROOT/'.work/reference')
reference=args.parse_args().reference
OUT=ROOT/'项目文档'; OUT.mkdir(exist_ok=True)
FIG=ROOT/'docs/diagrams'
figs={x['id']:x['title'] for x in json.loads((FIG/'manifest.json').read_text(encoding='utf-8'))}
figs.update({'prototype-login':'登录页面原型','prototype-chat':'智能问答页面原型','prototype-history':'历史会话页面原型',
 'prototype-faq':'常见问题入口原型','prototype-kb':'知识库文档管理页面原型','prototype-users':'用户与权限管理页面原型','prototype-stats':'问答统计页面原型'})

def after(p,text='',style=None):
    elem=OxmlElement('w:p'); p._p.addnext(elem)
    n=Paragraph(elem,p._parent)
    if style: n.style=style
    n.add_run(text)
    return n

def picture_after(p,ident,number):
    p.paragraph_format.keep_with_next=True
    pic=after(p)
    pic.paragraph_format.keep_with_next=True
    pic.paragraph_format.line_spacing=1
    pic.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=pic.add_run(); shape=run.add_picture(str(FIG/(ident+'.png')),width=Cm(16))
    if shape.height>Cm(14.8):
        ratio=Cm(14.8)/shape.height; shape.width=int(shape.width*ratio); shape.height=Cm(14.8)
    shape._inline.docPr.set('descr',figs[ident])
    caption=after(pic,f'图 {number}  {figs[ident]}','Caption')
    caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if ident.startswith('prototype'):
        note=after(caption,'页面采用当前前端和示例数据展示交互结构，图中数值不代表实测结果。')
        note.runs[0].font.size=Pt(9)
    return caption

def replace_table(doc,old,rows,widths=None):
    t=doc.add_table(rows=0,cols=len(rows[0]))
    for row in rows:
        cells=t.add_row().cells
        for c,text in zip(cells,row): c.text=str(text)
    old._tbl.addprevious(t._tbl); old._tbl.getparent().remove(old._tbl)
    if widths:
        t.autofit=False
        for row in t.rows:
            for cell,w in zip(row.cells,widths): cell.width=Cm(w)
    return t

COMMON={
 'QaController 调用 QaService 开始处理。':'MessageController 通过 MessageService 校验会话并保存问题，再调用 RagChatService。',
 'QaService 调用 Retriever 执行向量化与 Milvus 检索，得到候选片段。':'RagChatService 调用 RagRetrieveService 执行向量检索，再复核来源的当前权限和状态。',
 'QaService 调用 Reranker 对候选片段重排。':'RagChatService 调用 rerank 方法，按相似度取前 3 个可靠片段。',
 '限制文件类型与大小，对上传文件进行校验与消毒，防止恶意脚本。':'校验文件扩展名、大小和文件名路径，文件由专用解析器处理；病毒扫描与内容消毒为部署增强项。',
 '限制文件类型与大小，对上传播文件进行校验与消毒，防止恶意脚本。':'校验文件扩展名、大小和文件名路径，文件由专用解析器处理；病毒扫描与内容消毒为部署增强项。',
 '系统将企业内部散落在 Word、PDF、Excel、Markdown 与内部 Wiki 等各类文档中的资料进行统一接入、解析、数据清洗、文本分块与向量化处理，并保存到向量数据库 Milvus 中。':'系统统一接入 PDF、DOCX、TXT、Markdown 文档，完成解析、清洗、分块与向量化后保存到 Milvus。Excel 与内部 Wiki 接入属于后续增强。',
 'RAG 服务完成 Query Rewrite、Embedding 向量化、Metadata Filter 过滤、Milvus 向量检索、Top-K 召回、相似度阈值过滤、Rerank 重排、Context 构建':'RAG 服务完成 Embedding 向量化、Metadata Filter 过滤、Milvus Top-K 检索、相似度阈值及当前资源权限复核、相似度排序与 Context 构建',
 '相似度阈值过滤与 Rerank 重排':'相似度阈值过滤与排序截取；独立重排模型为后续增强',
 '6. 依据相似度阈值过滤低相关结果，必要时执行 Rerank 重排。':'6. 过滤低于阈值或当前无权访问的结果，按相似度取前 3 个可靠片段。',
 '支持部门管理与角色管理，内置普通员工、部门管理员、系统管理员三类角色。':'提供部门与角色列表、用户归属调整，内置普通员工、部门管理员、系统管理员三类角色；部门与角色增删改为后续增强。',
 '系统管理员开通用户账号并分配角色，用户可归属一个或多个部门。':'演示账号由初始化脚本创建，系统管理员可调整已有账号的角色和部门；每个账号归属一个部门。',
 '2. 选择 PDF/DOCX/TXT/MD 格式文件，填写分类与权限等级。':'2. 选择 PDF/DOCX/TXT/MD 格式文件，初始部门与权限继承知识库，后续可单独调整。',
 '1a. 知识库创建时未设置部门，系统默认为公开知识库。':'1a. 系统管理员可创建全司知识库，并明确选择访问级别；部门管理员必须选择本部门。',
 '·DeepSeek 等开源大语言模型的接口文档：用于智能回答功能的接口设计参考。':'·模型服务的 OpenAI 兼容接口文档：用于智能回答功能的接口设计参考。',
 '接口提供商：DeepSeek 等开源大模型 API。':'接口提供商：本地 LM Studio 的 OpenAI 兼容接口，模型可通过环境变量切换。',
 'DocumentService：文档服务类，负责文档上传、解析、清洗与状态管理。':'DocumentService：校验资源权限、保存上传文件与文档记录，触发异步管道并提供查询、重处理和删除操作。',
 'DocumentService 调用 TextParser 解析文档，返回纯文本内容。':'DocumentService 保存文件和 PARSING 记录后返回文档 ID，异步触发 DocumentPipelineService。',
 'DocumentService 调用 DataCleaner 清洗文本。':'DocumentPipelineService 通过 Tika 或 UTF-8 文件读取器提取文本，再调用 TextCleaner 清洗。',
 'DocumentService 调用 Chunker 进行分块，得到 Chunk 列表。':'DocumentPipelineService 调用 Chunker 分块，附加文档、部门和权限等元数据。',
 'DocumentService 调用 EmbeddingService 将 Chunk 向量化。':'VectorStore.add 内部调用 EmbeddingModel 生成向量。',
 'DocumentService 调用 MilvusRepository 将向量写入 Milvus，并返回入库结果。':'MilvusVectorStore 写入向量，DocumentPipelineService 将记录更新为 READY；异常时记录 FAILED 并按策略重试。',
 'DocumentService 更新文档状态并向管理员返回处理结果。':'管理页面轮询文档详情，展示最新处理状态和分块数量。',
 'Java 17':'Java 21','JDK 17':'JDK 21','MySQL 5.7+':'MySQL 8.4',
 '/api/qa/ask':'/api/messages','/api/kb/document/upload':'/api/kb/{kbId}/documents','/api/chat/history':'/api/conversations',
 '本项目使用 DeepSeek':'本项目默认使用本地 Gemma 模型，可通过配置切换',
 '（如 DeepSeek）':'（默认 Gemma，可通过配置切换）',
 '服务器端：操作系统 Linux（CentOS 7+ / Ubuntu 18.04+）或 Windows Server；后端环境 Java 21 + Spring Boot；数据库 MySQL 8.4；向量数据库 Milvus 2.x；Web 服务器 Nginx 1.16+，用于反向代理与静态资源服务。':
 '服务器端：支持 Windows 或 Linux，后端使用 Java 21、Spring Boot 4.0.8、Spring AI 2.0.0，业务数据库 MySQL 8.4、向量数据库 Milvus 2.6。生产环境应通过反向代理配置 HTTPS。',
 '支持流式输出以避免等待过长':'当前采用完整响应与等待提示；流式输出列为后续增强',
 '支持流式渲染回答':'当前渲染完整回答并显示加载状态；流式渲染为后续增强',
 '对历史会话进行分页加载，减少接口压力。':'当前加载本人会话列表；数据量增长后增加分页与按需加载。',
 '会话列表分页展示，避免一次性加载过多数据。':'当前返回本人会话列表；分页展示列为后续增强，以控制大数据量开销。',
 '按时间倒序分页返回会话及其中的消息记录。':'会话列表按 ID 倒序返回；GET /api/conversations/{id}/messages 按消息 ID 升序返回本人会话消息。',
 '系统按时间倒序展示会话中的问题与回答。':'系统按时间升序展示会话中的问题与回答。',
 '对问题做 Query Rewrite 改写，规范表达。':'当前直接使用原问题检索；Query Rewrite 改写列为后续增强。',
 '对问题执行 Query Rewrite 改写。':'当前直接检索原问题；独立 Query Rewrite 模块列为后续增强。',
 '系统对问题进行 Query Rewrite 规范化改写。':'当前直接使用原问题检索；问题规范化改写为后续增强。',
 '将改写后的问题文本向量化。':'将问题文本向量化。',
 '对候选结果进行 Rerank 重排。':'按相似度排序并截取前 3 个可靠片段；独立重排模型为后续增强。',
 '对候选结果执行 Rerank 重排，精选最相关的片段。':'按相似度排序并截取前 3 个可靠片段。独立重排模型为后续增强。',
 '用户可归属一个或多个部门':'当前用户归属一个部门',
 '普通员工仅可访问公开资料':'普通员工可访问全司公开资料及本部门允许的资料',
 '敏感配置与大模型密钥使用安全的密钥管理机制进行加密存储。':'数据库密码、JWT 密钥和模型凭据通过环境变量配置，本地 .env 不入库；生产部署应接入密钥管理服务。',
 '敏感配置与大模型密钥采用加密方式管理。':'敏感配置由环境变量注入，生产环境应使用专门的密钥管理服务。',
 '密钥管理：大模型与向量数据库的连接密钥采用集中配置并加密保存，定期轮换。':'密钥管理：连接凭据通过环境变量注入；部署时应限制访问并建立轮换机制。',
 '对用户输入进行长度与敏感词预处理。':'前后端校验问题非空且不超过 2000 字；当前不提供独立敏感词过滤器。',
 '限制大模型仅依据知识库内容回答，对用户输入进行敏感词过滤，防止越权诱导与注入。':'通过系统提示要求依据检索片段回答，并在检索和生成前执行资源权限检查。提示词无法单独保证抵御所有注入，应持续开展对抗测试。',
 '支持用户注册与登录':'支持已有企业账号登录',
 '支持对用户问题的 Query Rewrite 改写，提高检索命中率。':'问题改写为后续增强；当前直接使用原问题检索。',
 '支持对用户问题的 Query Rewrite 改写，提升检索命中率。':'问题改写为后续增强；当前直接使用原问题检索。',
 '支持 Rerank 重排，精选最相关片段并构建高质量上下文。':'当前按相似度选取前 3 个可靠片段构建上下文；独立重排模型为后续增强。',
 '入参：问题文本、会话ID、部门ID。':'入参：content 问题文本、conversationId 本人会话 ID；部门与身份从 JWT 对应账号获取。',
 '出参：回答内容、引用来源列表、会话ID。':'出参：messageId、content、sources、answered。',
 '入参：文档文件、知识库ID、部门、分类、权限等级。':'入参：路径 kbId、multipart 文件 file；默认继承知识库部门与权限。',
 '入参：用户ID、页码、每页条数。':'入参：Authorization 请求头；查询消息时提供路径会话 ID。',
 '出参：会话列表、总条数。':'出参：本人会话列表；消息通过 /api/conversations/{id}/messages 查询。',
 '出参：访问凭证、用户ID、角色。':'出参：token、user 用户信息。统一外层结构为 code、msg、data。',
 '支持 Rerank 重排，精选最相关内容。':'当前按相似度排序并截取前 3 个可靠片段；独立重排模型为后续增强。',
 'Word、PDF、Excel、Markdown 与内部 Wiki 等各类文档中的资料':'PDF、DOCX、TXT、Markdown 文档中的资料；Excel 与内部 Wiki 接入属于后续增强。对上述直接支持的文档',
 'Word、PDF、Excel、Markdown、内部 Wiki 等各类文档中的资料':'PDF、DOCX、TXT、Markdown 文档中的资料',
 '系统管理员创建部门、角色与权限资源。':'当前内置部门与三类角色；系统管理员可查看部门与角色列表、修改用户归属与账号状态。部门及角色的增删改列为后续增强。',
 '1. 管理员创建/编辑部门信息。':'1. 管理员查看部门信息；部门增删改为后续增强。',
 '2. 管理员创建/编辑角色并分配权限资源。':'2. 管理员查看内置角色；动态角色配置为后续增强。',
 '3. 管理员开通用户账号，归属部门并分配角色。':'3. 当前通过初始化脚本创建演示账号，管理员为已有账号调整部门和角色；账号开通界面为后续增强。',
 '统计数据需支持按部门与时间范围筛选。':'当前统计支持时间范围选择及部门分布展示；按指定部门筛选为后续增强。',
 '前端遵循 ESLint 规范':'前端通过 TypeScript 类型检查',
 '对密码进行加盐哈希校验，防止明文比对导致的泄露。':'根据用户名查询账号，再使用 BCrypt 校验密码哈希。',
 '将 token 返回前端，并建立会话记录。':'将 token 和用户信息返回前端；聊天会话由 POST /api/conversations 单独创建。',
}

def transform(text):
    for a,b in COMMON.items(): text=text.replace(a,b)
    return text

def normalize(doc):
    for sec in doc.sections:
        sec.page_width=Cm(21);sec.page_height=Cm(29.7)
        sec.top_margin=sec.bottom_margin=Cm(1.9); sec.left_margin=sec.right_margin=Cm(2.2)
        sec.header_distance=sec.footer_distance=Cm(.8)
        for grid in sec._sectPr.findall(qn('w:docGrid')):sec._sectPr.remove(grid)
        for border in sec._sectPr.findall(qn('w:pgBorders')):sec._sectPr.remove(border)
    for name,size in [('Normal',10.5),('Title',24),('Heading 1',16),('Heading 2',13),('Heading 3',11.5),('Heading 4',11),('Caption',9.5)]:
        st=doc.styles[name]; st.font.name='Microsoft YaHei'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
        pf=st.paragraph_format;pf.space_before=Pt(10 if 'Heading' in name else 0);pf.space_after=Pt(5);pf.line_spacing=1.25
        if name.startswith('Heading'): pf.keep_with_next=True
        for border in st.element.xpath('.//w:pBdr'): border.getparent().remove(border)
    for p in doc.paragraphs:
        pp=p._p.get_or_add_pPr()
        for tag in ('spacing','snapToGrid','contextualSpacing','textAlignment'):
            for node in pp.findall(qn('w:'+tag)):pp.remove(node)
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp.append(snap)
        pf=p.paragraph_format; pf.left_indent=pf.right_indent=pf.first_line_indent=Pt(0)
        if p.text.startswith('表5-'):pf.keep_with_next=True
        pf.line_spacing=1.2; pf.space_before=Pt(0); pf.space_after=Pt(5)
        pf.keep_together=False
        pf.widow_control=True
        for b in p._p.xpath('.//w:pBdr'): b.getparent().remove(b)
        for r in p.runs:
            r.font.name='Microsoft YaHei';r.font.color.rgb=RGBColor(0,0,0)
            r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
            if not p._p.xpath('.//w:drawing'): r.font.size=None
        if p.style.name.startswith('Heading'):
            pf.keep_with_next=True; pf.space_before=Pt(12)
            p.text=re.sub(r'[、：:]+',' ',p.text).strip()
    for t in doc.tables:
        t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
        n=len(t.columns)
        widths={2:[3.4,13],3:[3.4,4.0,9],4:[2.2,3.6,2.5,8.1],5:[1.2,3.8,2.8,2,6.6]}.get(n,[16.4/n]*n)
        for col,width in zip(t.columns,widths):col.width=Cm(width)
        for node in t._tbl.tblPr.findall(qn('w:tblInd')):t._tbl.tblPr.remove(node)
        tw=t._tbl.tblPr.find(qn('w:tblW'));tw.set(qn('w:w'),str(int(Cm(16.4).twips)));tw.set(qn('w:type'),'dxa')
        for node in t._tbl.tblPr.findall(qn('w:tblBorders')):t._tbl.tblPr.remove(node)
        borders=OxmlElement('w:tblBorders')
        for tag in ('top','left','bottom','right','insideH','insideV'):
            e=OxmlElement('w:'+tag);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        t._tbl.tblPr.append(borders)
        for i,row in enumerate(t.rows):
            row.height=None
            for h in row._tr.xpath('.//w:trHeight'):h.getparent().remove(h)
            trPr=row._tr.get_or_add_trPr(); cs=OxmlElement('w:cantSplit');trPr.append(cs)
            if i==0: trPr.append(OxmlElement('w:tblHeader'))
            for col_index,cell in enumerate(row.cells):
                cell.width=Cm(widths[min(col_index,len(widths)-1)])
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp=cell._tc.get_or_add_tcPr()
                for node in cp.findall(qn('w:tcBorders')):cp.remove(node)
                cell_borders=OxmlElement('w:tcBorders')
                for tag in ('top','left','bottom','right'):
                    e=OxmlElement('w:'+tag);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');cell_borders.append(e)
                cp.append(cell_borders)
                for shd in cp.findall(qn('w:shd')):cp.remove(shd)
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E6EEF6' if i==0 else 'FFFFFF');cp.append(shd)
                margins=OxmlElement('w:tcMar')
                for tag in ('top','bottom','left','right'):
                    e=OxmlElement('w:'+tag);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
                cp.append(margins)
                for p in cell.paragraphs:
                    pp=p._p.get_or_add_pPr()
                    for tag in ('spacing','snapToGrid','ind'):
                        for node in pp.findall(qn('w:'+tag)):pp.remove(node)
                    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp.append(snap)
                    p.paragraph_format.line_spacing=1.12;p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                    p.paragraph_format.keep_with_next=False
                    for r in p.runs:
                        r.font.name='Microsoft YaHei';r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0);r.font.bold=i==0
                        r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    # Simple page numbers, no old project header text or ornamental rules.
    for section in doc.sections:
        for part in (section.header,section.first_page_header,section.even_page_header,section.footer,section.first_page_footer,section.even_page_footer):
            for node in list(part._element):part._element.remove(node)
            part.add_paragraph()
        footer=section.footer.paragraphs[0];footer.clear();footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for node in footer._p.xpath('.//w:pBdr'):node.getparent().remove(node)
        run=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
    for style in doc.styles:
        for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)

for kind,filename in [('requirements','系统需求分析说明书'),('design','系统设计说明书')]:
    doc=Document(reference/(filename+'.docx'))
    for style_name in ['Title','Caption','Normal','Heading 1','Heading 2','Heading 3','Heading 4']:
        if style_name not in doc.styles:
            doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    original_tables=list(doc.tables)
    # All old pictures are from unrelated projects; remove relationships too.
    for node in doc._element.xpath('.//w:drawing | .//w:pict | .//w:object'):
        node.getparent().remove(node)
    for rid,rel in list(doc.part.rels.items()):
        if rel.reltype.endswith('/image'):doc.part.drop_rel(rid)
    for p in list(doc.paragraphs):
        if not p.text.strip():p._p.getparent().remove(p._p)
        else:p.text=transform(p.text)
    for table in original_tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:p.text=transform(p.text)
    # Keep student/team details; remove empty assessment/signature template sheet.
    original_tables[1]._tbl.getparent().remove(original_tables[1]._tbl)
    title=doc.paragraphs[0];title.text='企业内部知识库客服系统\n'+filename;title.style='Title'
    title.paragraph_format.space_before=Pt(24);title.paragraph_format.space_after=Pt(20)
    doc.core_properties.title='企业内部知识库客服系统 '+filename
    doc.core_properties.author='Ai大坝组';doc.core_properties.last_modified_by='Ai大坝组'
    doc.core_properties.comments=''
    overview=next(p for p in doc.paragraphs if p.style.name.startswith('Heading'))
    overview.paragraph_format.page_break_before=True
    scope='本版本以文档入库、权限检索、带来源问答、个人会话、评价和基础统计为核心。当前采用完整响应、相似度排序和内置角色；独立问题改写、模型重排、流式输出、历史分页与动态角色管理属于后续增强。性能数值为验收目标，须在指定模型与硬件上测试，不能视为已通过的测量结果。'
    after(title,'项目组  Ai大坝组    版本日期  2026年9月18日')
    after(overview,scope)
    count=0
    def insert(p,key):
        global count
        count+=1;picture_after(p,key,f'{"需" if kind=="requirements" else "设"}-{count:02d}')
    if kind=='requirements':
        mapping={'员工端子系统用例图':'usecase-employee','管理端子系统用例图':'usecase-admin','登录页原型图':'prototype-login','问答页原型图':'prototype-chat','历史会话页原型图':'prototype-history','常见问题入口原型图':'prototype-faq','知识库管理页原型图':'prototype-kb','用户、部门、角色管理原型图':'prototype-users','统计与分析页原型图':'prototype-stats','智能问答活动图':'flow-qa','文档入库活动图':'flow-ingest'}
        for p in list(doc.paragraphs):
            if '此处插入' in p.text:
                ident=next(key for text,key in mapping.items() if text in p.text)
                p.text='';insert(p,ident)
            elif p.text.startswith('2.4 系统整体结构'):insert(p,'architecture')
            elif p.text.startswith('3.2 系统主 Use Case'):insert(p,'usecase-main')
    else:
        mapping={'2功能模块设计':'modules','4.1.1 入库流程':'flow-ingest','4.1.2 类图说明':'class-ingest','4.1.3 时序图说明':'sequence-ingest',
            '4.2.1 问答流程':'flow-qa','4.2.2 类图说明':'class-chat','4.2.3 时序图说明':'sequence-chat','4.3用户权限管理流程':'access','4.4知识问答页面设计':'prototype-chat','5数据库设计':'er'}
        for p in list(doc.paragraphs):
            if p.text in mapping:insert(p,mapping[p.text])
        # Actual API contract: identity is taken from JWT, never from request departmentId/userId.
        api_tables=[
         [['位置','字段','类型','说明'],['body','username','String','用户名'],['body','password','String','密码']],
         [['位置','字段','类型','说明'],['data','token','String','JWT 访问令牌'],['data','user','UserInfoVO','用户 ID、部门与角色信息']],
         [['位置','字段','类型','说明'],['body','content','String','问题正文，1 至 2000 字'],['body','conversationId','Long','本人会话 ID']],
         [['位置','字段','类型','说明'],['data','messageId','Long','回答消息 ID'],['data','content','String','回答或兜底提示'],['data','sources','List','documentId/title/chunkId/similarity'],['data','answered','Boolean','是否生成有效回答']],
         [['位置','字段','类型','说明'],['path','kbId','Long','知识库 ID'],['multipart','file','File','PDF/DOCX/TXT/MD，最多 10 MB']],
         [['位置','字段','类型','说明'],['data','id','Long','文档 ID'],['data','status','String','PARSING/READY/FAILED'],['data','chunkCount','Integer','切片数量']],
         [['位置','字段','类型','说明'],['header','Authorization','String','Bearer JWT，身份由服务端解析'],['path','id','Long','查询 /conversations/{id}/messages 时提供']],
         [['位置','字段','类型','说明'],['data','会话列表','List','id/title/createdAt，按本人隔离'],['data','消息列表','List','按会话 ID 查询，包含 content/sources/answered']]]
        for old,rows in zip(original_tables[2:10],api_tables):replace_table(doc,old,rows,[2,3.6,2.5,8.3])
        sql=(ROOT/'backend/src/main/resources/db/schema.sql').read_text(encoding='utf-8')
        sql=re.sub(r'--[^\n]*','',sql)
        for old,name in zip(original_tables[10:],['user','department','role','permission','knowledge_base','document','conversation','message','feedback']):
            body=re.search(r'create table if not exists '+name+r'\s*\((.*?)\n\);',sql,re.S).group(1)
            rows=[['字段','类型','约束与说明']]
            for field in re.split(r',\s*(?![^()]*\))',body):
                field=' '.join(field.split())
                match=re.match(r'(\w+)\s+(bigint|varchar\(\d+\)|tinyint|datetime|double|text|int|date)(.*)',field)
                if match:
                    col,typ,tail=match.groups()
                    desc='主键 自增' if 'primary key' in tail else ('必填' if 'not null' in tail else '可空')
                    if 'default' in tail:desc+='；'+tail[tail.index('default'):]
                    rows.append([col,typ,desc])
            replace_table(doc,old,rows,[4.6,3.2,8.6])
        for p in doc.paragraphs:
            t=p.text
            if t.startswith('TextParser：'):p.text='TikaDocumentReader：用于 PDF、DOCX 文本提取；MD/TXT 由文件读取器按 UTF-8 读取。'
            elif t.startswith('DataCleaner：'):p.text='TextCleaner：清洗文本中的多余空白与重复段落。'
            elif t.startswith('EmbeddingService：'):p.text='EmbeddingModel：由 Spring AI VectorStore 在添加分块或检索时调用。'
            elif t.startswith('MilvusRepository：'):p.text='MilvusVectorStore 与 MilvusSchemaInitializer：负责向量写入、检索、分块查询与删除。'
            elif t.startswith('各服务之间的依赖关系：'):p.text='DocumentService 校验权限并保存文档；DocumentPipelineService 异步解析、清洗与分块，调用 VectorStore 写入，并更新文档状态。'
            elif t.startswith('QaController：'):p.text='MessageController：接收问题、校验输入，调用 MessageService。'
            elif t.startswith('QaService：'):p.text='MessageService 与 RagChatService：分别负责会话持久化和检索增强生成编排。'
            elif t.startswith('QueryRewriter：'):p.text='问题改写模块：后续增强接口，当前版本直接使用原问题检索。'
            elif t.startswith('Retriever：'):p.text='RagRetrieveService：使用 VectorStore 检索并排序；RagChatService 执行阈值与来源权限复核。'
            elif t.startswith('Reranker：'):p.text='rerank 方法：当前按相似度截取前 N 块，未调用独立重排模型。'
            elif t.startswith('LlmClient：'):p.text='ChatClient：Spring AI 聊天客户端，使用可配置的本地模型接口。'
            elif t.startswith('流程依赖：'):p.text='MessageController → MessageService → RagChatService → RagRetrieveService / ResourceAccessService / ChatClient。'
            elif 'QaService 调用 QueryRewriter' in t:p.text='RagChatService 使用原问题开始检索，历史消息作为生成上下文。'
            else:
                for a,b in [('QaController','MessageController'),('QaService','RagChatService'),('TextParser','TikaDocumentReader'),('DataCleaner','TextCleaner'),('EmbeddingService','EmbeddingModel'),('MilvusRepository','MilvusVectorStore'),('Reranker','RagRetrieveService'),('LlmClient','ChatClient')]:t=t.replace(a,b)
                if not p._p.xpath('.//w:drawing'):p.text=t
        # The replacement loop above must preserve figure paragraphs (do not assign text there).
    normalize(doc)
    # Remove manual page breaks inherited from the old template, excluding the cover boundary.
    for p in doc.paragraphs:
        for b in p._p.xpath('.//w:br[@w:type="page"]'):b.getparent().remove(b)
    for p in doc.paragraphs:
        if p.text.startswith('图 '):p.paragraph_format.keep_with_next=False
    path=OUT/(filename+'.docx');doc.save(path)
    print(filename,count,'figures',len(doc.tables),'tables')
